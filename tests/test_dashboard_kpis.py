"""Known-total reconciliation (Astra's publication checklist, item 1):
the generator's independently-tracked truth (`known_totals.json`) must
match `kpis.compute_kpis`'s recomputation from the CSVs on disk, and both
must match the literal values written into the built dashboard workbooks.
"""

from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from team_dashboard.dashboard import build_dashboards
from team_dashboard.kpis import compute_kpis
from team_dashboard.teams import TEAM_NAMES

AS_OF = date(2026, 9, 21)


def test_kpis_match_known_totals_per_team(data_dir, known_totals):
    result = compute_kpis(data_dir, AS_OF)
    for team in TEAM_NAMES:
        truth = known_totals["teams"][team]
        computed = result["teams"][team]
        assert computed.planned_total == truth["planned_total"], team
        assert computed.done_total == truth["done_total"], team
        assert computed.overdue_total == truth["overdue_total"], team
        assert computed.open_blockers_total == truth["open_blockers_total"], team
        assert computed.open_blockers_by_owner == truth["open_blockers_by_owner"], team
        assert computed.achievements_this_week == truth["achievements_this_week"], team
        assert computed.weekly_done_trend == truth["weekly_done_trend"], team


def test_kpis_match_known_totals_overall(data_dir, known_totals):
    result = compute_kpis(data_dir, AS_OF)
    overall = result["overall"]
    truth = known_totals["overall"]
    assert overall.planned_total == truth["planned_total"]
    assert overall.done_total == truth["done_total"]
    assert overall.overdue_total == truth["overdue_total"]
    assert overall.open_blockers_total == truth["open_blockers_total"]
    assert overall.achievements_this_week == truth["achievements_this_week"]
    assert overall.weekly_done_trend == truth["weekly_done_trend"]


def test_master_dashboard_kpi_table_matches_known_totals(data_dir, known_totals, tmp_path):
    result = build_dashboards(data_dir, tmp_path / "dashboards", as_of=AS_OF)
    wb = load_workbook(result["master"])
    ws = wb["Overview"]

    header = [c.value for c in ws[4]]
    assert header == ["Team", "Planned", "Done", "Done %", "Overdue", "Open Blockers", "Achievements This Week"]

    rows_by_team = {}
    for row in ws.iter_rows(min_row=5, max_row=4 + len(TEAM_NAMES) + 1, values_only=True):
        if row[0] is None:
            break
        rows_by_team[row[0]] = row

    for team in TEAM_NAMES:
        truth = known_totals["teams"][team]
        row = rows_by_team[team]
        assert row[1] == truth["planned_total"]
        assert row[2] == truth["done_total"]
        assert row[4] == truth["overdue_total"]
        assert row[5] == truth["open_blockers_total"]
        assert row[6] == truth["achievements_this_week"]

    overall_row = rows_by_team["All teams"]
    truth = known_totals["overall"]
    assert overall_row[1] == truth["planned_total"]
    assert overall_row[2] == truth["done_total"]
    assert overall_row[4] == truth["overdue_total"]
    assert overall_row[5] == truth["open_blockers_total"]
    assert overall_row[6] == truth["achievements_this_week"]


def test_team_dashboard_scoped_to_one_team(data_dir, known_totals, tmp_path):
    result = build_dashboards(data_dir, tmp_path / "dashboards", as_of=AS_OF)
    team = "Team Aurora"
    wb = load_workbook(result["teams"][team])
    ws = wb["Overview"]
    row = list(ws.iter_rows(min_row=5, max_row=5, values_only=True))[0]
    truth = known_totals["teams"][team]
    assert row[0] == team
    assert row[1] == truth["planned_total"]
    assert row[2] == truth["done_total"]
    # A team copy has no "All teams" rollup row.
    next_row = list(ws.iter_rows(min_row=6, max_row=6, values_only=True))[0]
    assert next_row[0] != "All teams"


def test_master_dashboard_has_a_native_chart(data_dir, tmp_path):
    result = build_dashboards(data_dir, tmp_path / "dashboards", as_of=AS_OF)
    wb = load_workbook(result["master"])
    ws = wb["Overview"]
    assert len(ws._charts) == 1


def test_overdue_and_blocker_detail_sheets_match_kpi_counts(data_dir, known_totals, tmp_path):
    result = build_dashboards(data_dir, tmp_path / "dashboards", as_of=AS_OF)
    wb = load_workbook(result["master"])

    overdue_ws = wb["Overdue Tasks"]
    overdue_rows = list(overdue_ws.iter_rows(min_row=2, values_only=True))
    overdue_rows = [r for r in overdue_rows if r[0] is not None]
    assert len(overdue_rows) == known_totals["overall"]["overdue_total"]

    blockers_ws = wb["Open Blockers"]
    blocker_rows = list(blockers_ws.iter_rows(min_row=2, values_only=True))
    blocker_rows = [r for r in blocker_rows if r[0] is not None]
    assert len(blocker_rows) == known_totals["overall"]["open_blockers_total"]
