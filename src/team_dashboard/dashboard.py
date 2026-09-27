"""Builds the master dashboard workbook and the distributed per-team copies.

Both are built the same way -- computed from the same CSVs via `kpis.py` --
the team copies just scope the KPI table, detail sheets, and trend chart to
one team. All KPI figures are written as plain values (not live formulas):
this is the Python "refresh" path described in the README, not a workbook
with a live Power Query connection (which needs Excel to author; see
`powerquery/README.md`).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .kpis import TeamKPIs, compute_kpis
from .teams import TEAM_NAMES

KPI_HEADERS = ["Team", "Planned", "Done", "Done %", "Overdue", "Open Blockers", "Achievements This Week"]


def _slug(team: str) -> str:
    return team.lower().removeprefix("team ").replace(" ", "-")


def _kpi_row(team: str, tk: TeamKPIs) -> list:
    done_pct = round(100 * tk.done_total / tk.planned_total, 1) if tk.planned_total else 0.0
    return [team, tk.planned_total, tk.done_total, done_pct, tk.overdue_total, tk.open_blockers_total, tk.achievements_this_week]


# Leading characters that would make Excel/openpyxl treat a plain-text
# value as an executable formula (`=`, `+`, `-`, `@`) or that spreadsheet
# software can otherwise misinterpret (tab, CR). Source text -- task
# titles, descriptions, summaries -- is untrusted input that reached this
# dashboard via `sync`'s CSVs, and must render as inert text, never as a
# formula `sync` or a human never asked for.
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r")


def _write_cell(ws, row: int, col: int, val):
    cell = ws.cell(row=row, column=col, value=val)
    if isinstance(val, str) and val.startswith(_FORMULA_TRIGGER_CHARS):
        # openpyxl auto-detects a leading "=" as a formula and sets
        # data_type "f" accordingly; forcing it back to "s" (string) here
        # writes the exact same text back out as an inert value -- Excel
        # never recalculates it, it just displays literally.
        cell.data_type = "s"
    return cell


def _write_table(ws, start_row: int, headers: list[str], rows: list[list]) -> int:
    for c, h in enumerate(headers, start=1):
        cell = _write_cell(ws, start_row, c, h)
        cell.font = Font(bold=True)
    for r, row in enumerate(rows, start=start_row + 1):
        for c, val in enumerate(row, start=1):
            _write_cell(ws, r, c, val)
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, len(str(h)) + 2)
    return start_row + 1 + len(rows)  # next free row


def _build_overview(ws, teams: list[str], kpi_result: dict, as_of: date):
    ws["A1"] = "Team progress dashboard"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"As of: {as_of.isoformat()}"

    kpi_rows = [_kpi_row(t, kpi_result["teams"][t]) for t in teams]
    if len(teams) > 1:
        overall = kpi_result["overall"]
        kpi_rows.append(_kpi_row("All teams", overall))
    next_row = _write_table(ws, 4, KPI_HEADERS, kpi_rows)

    # Weekly done trend table (source data for the native chart)
    trend_start = next_row + 2
    ws.cell(row=trend_start, column=1, value="Weekly done trend").font = Font(bold=True)
    header_row = trend_start + 1
    week_endings = kpi_result["week_endings"]
    trend_headers = ["Week Ending"] + teams
    trend_rows = []
    for w in week_endings:
        wk = w.isoformat()
        row = [w] + [kpi_result["teams"][t].weekly_done_trend[wk] for t in teams]
        trend_rows.append(row)
    _write_table(ws, header_row, trend_headers, trend_rows)

    chart = LineChart()
    chart.title = "Tasks done per week"
    chart.y_axis.title = "Tasks done"
    chart.x_axis.title = "Week ending"
    data_ref = Reference(ws, min_col=2, max_col=1 + len(teams), min_row=header_row, max_row=header_row + len(trend_rows))
    cats_ref = Reference(ws, min_col=1, max_col=1, min_row=header_row + 1, max_row=header_row + len(trend_rows))
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, f"A{header_row + len(trend_rows) + 3}")

    for c in range(1, len(trend_headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(14, ws.column_dimensions[get_column_letter(c)].width or 0)


def _build_detail_sheets(wb: Workbook, teams: list[str], kpi_result: dict):
    overdue_headers = ["task_id", "team", "owner", "title", "status", "planned_date", "days_overdue"]
    overdue_rows = []
    for t in teams:
        for row in kpi_result["teams"][t].overdue_tasks:
            days = (kpi_result["as_of"] - date.fromisoformat(row["planned_date"])).days
            overdue_rows.append([row["task_id"], row["team"], row["owner"], row["title"], row["status"], row["planned_date"], days])
    ws = wb.create_sheet("Overdue Tasks")
    _write_table(ws, 1, overdue_headers, overdue_rows)

    blocker_headers = ["blocker_id", "team", "task_id", "owner", "description", "raised_date", "age_days"]
    blocker_rows = []
    for t in teams:
        for row in kpi_result["teams"][t].open_blockers:
            blocker_rows.append(
                [row["blocker_id"], row["team"], row["task_id"], row["owner"], row["description"], row["raised_date"], row["age_days"]]
            )
    blocker_rows.sort(key=lambda r: r[-1], reverse=True)
    ws = wb.create_sheet("Open Blockers")
    _write_table(ws, 1, blocker_headers, blocker_rows)


def _build_achievements_sheet(wb: Workbook, teams: list[str], data_dir: Path, as_of: date):
    from . import csv_io, schema

    rows = csv_io.read_rows(data_dir / schema.ACHIEVEMENTS.filename, schema.ACHIEVEMENTS)
    this_week = [r for r in rows if r["team"] in teams and r["week_ending"] == as_of.isoformat()]
    ws = wb.create_sheet("Achievements This Week")
    _write_table(ws, 1, ["achievement_id", "team", "member", "week_ending", "description"], [list(r.values()) for r in this_week])


def _build_workbook(teams: list[str], kpi_result: dict, data_dir: Path, as_of: date) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    _build_overview(ws, teams, kpi_result, as_of)
    _build_detail_sheets(wb, teams, kpi_result)
    _build_achievements_sheet(wb, teams, data_dir, as_of)
    return wb


def build_dashboards(data_dir: Path, out_dir: Path, as_of: date, teams: list[str] | None = None) -> dict:
    """Builds the master workbook and one distributed copy per team.

    Returns {"master": path, "teams": {team: path}} plus the computed
    kpi_result dict (under "kpis"), so callers/tests can reconcile figures
    against `known_totals.json` without re-reading the workbook.
    """
    team_names = teams if teams is not None else list(TEAM_NAMES)
    kpi_result = compute_kpis(data_dir, as_of, teams=team_names)
    kpi_result["as_of"] = as_of

    out_dir.mkdir(parents=True, exist_ok=True)
    master_path = out_dir / "master-dashboard.xlsx"
    master_wb = _build_workbook(team_names, kpi_result, data_dir, as_of)
    master_wb.save(master_path)

    team_paths = {}
    for team in team_names:
        team_wb = _build_workbook([team], kpi_result, data_dir, as_of)
        path = out_dir / f"team-{_slug(team)}-dashboard.xlsx"
        team_wb.save(path)
        team_paths[team] = path

    return {"master": master_path, "teams": team_paths, "kpis": kpi_result}
