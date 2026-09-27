"""Builds one filled-in, submitted-looking input workbook for the demo run:
a deliberate mix of valid rows and one row from every bad-input class, so
`run_demo.sh` and the test suite can show `sync` actually rejecting things
into the exceptions report instead of silently dropping them.

Everything here is fabricated, same as `generator.py`.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from . import csv_io, schema
from .input_workbook import build_input_workbook


def build_sample_inbox_workbook(path: Path, data_dir: Path, as_of_iso: str) -> None:
    tasks = csv_io.read_rows(data_dir / schema.TASKS.filename, schema.TASKS)
    an_existing_task_id = tasks[0]["task_id"] if tasks else "T-AUR-0001"

    build_input_workbook(path)
    wb = load_workbook(path)

    tasks_ws = wb["New Tasks"]
    tasks_rows = [
        # valid
        ["T-DEMO-0001", "Team Aurora", "Priya Nandan", "Write onboarding doc for new hires", "In Progress", "2026-09-20", "", "2026-09-15"],
        # missing_required_field (blank title)
        ["T-DEMO-0002", "Team Aurora", "Priya Nandan", "", "In Progress", "2026-09-20", "", "2026-09-15"],
        # invalid_enum (bad status)
        ["T-DEMO-0003", "Team Aurora", "Priya Nandan", "Tidy up the backlog", "Almost Done", "2026-09-20", "", "2026-09-15"],
        # invalid_date
        ["T-DEMO-0004", "Team Aurora", "Priya Nandan", "Fix the flaky test", "In Progress", "20-09-2026", "", "2026-09-15"],
        # unknown_team
        ["T-DEMO-0005", "Team Denali", "Priya Nandan", "Investigate outage", "In Progress", "2026-09-20", "", "2026-09-15"],
        # unknown_member
        ["T-DEMO-0006", "Team Aurora", "Not A Person", "Review pull requests", "In Progress", "2026-09-20", "", "2026-09-15"],
        # duplicate_key (same task_id as the first valid row above)
        ["T-DEMO-0001", "Team Aurora", "Priya Nandan", "Duplicate of the first row", "In Progress", "2026-09-20", "", "2026-09-15"],
    ]
    for r, row in enumerate(tasks_rows, start=2):
        for c, val in enumerate(row, start=1):
            tasks_ws.cell(row=r, column=c, value=val)

    blockers_ws = wb["New Blockers"]
    blockers_rows = [
        # valid
        ["B-DEMO-0001", "Team Aurora", an_existing_task_id, "Priya Nandan", "Waiting on API keys", "2026-09-19", "Open", ""],
        # unknown_task_id
        ["B-DEMO-0002", "Team Aurora", "T-DOES-NOT-EXIST", "Priya Nandan", "Blocked on nothing real", "2026-09-19", "Open", ""],
        # invalid_number is exercised on weekly_updates, not here
    ]
    for r, row in enumerate(blockers_rows, start=2):
        for c, val in enumerate(row, start=1):
            blockers_ws.cell(row=r, column=c, value=val)

    updates_ws = wb["New Weekly Updates"]
    updates_rows = [
        # valid
        ["U-DEMO-0001", "Team Aurora", "Priya Nandan", as_of_iso, "Wrapped up the sprint early.", 32.5],
        # invalid_number (negative hours)
        ["U-DEMO-0002", "Team Aurora", "Priya Nandan", as_of_iso, "Logged negative hours by mistake.", -4],
    ]
    for r, row in enumerate(updates_rows, start=2):
        for c, val in enumerate(row, start=1):
            updates_ws.cell(row=r, column=c, value=val)

    achievements_ws = wb["New Achievements"]
    achievements_rows = [
        ["A-DEMO-0001", "Team Aurora", "Priya Nandan", as_of_iso, "Presented the roadmap to stakeholders."],
    ]
    for r, row in enumerate(achievements_rows, start=2):
        for c, val in enumerate(row, start=1):
            achievements_ws.cell(row=r, column=c, value=val)

    wb.save(path)
